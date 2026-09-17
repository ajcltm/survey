"""
오픈마켓 체감도조사 - 쿠팡 '중개서비스' 입점업체 판매자정보 수집기 (v3: 국내 입점업체 한정)
------------------------------------------------------------------------
2026-09-17 실제 쿠팡 페이지 구조를 확인하고 맞춘 버전입니다.

[검색결과 카드의 배송 배지 → 거래형태]
  logo_rocket_filter   : 로켓배송  → 판매자 '쿠팡(주)' = 직매입        → 제외
  logo_jikgu           : 로켓직구  → 해외 입점업체 판매(실제 확인: 중국 법인) → 제외 (국내 입점업체 한정)
  logo_rocket_merchant : 판매자로켓 → 입점업체 판매 + 쿠팡 물류(로켓그로스) → 중개 (기본 포함)
  배지 없음             : 판매자배송 → 입점업체 판매·직접 배송            → 중개 (포함)
  ※ 배지는 글자가 아닌 이미지라서 이미지 파일명으로 구분합니다.
  ※ 카드 단계는 1차 거름망이고, 상품페이지에서 한 번 더 거릅니다.
     - 판매자가 '쿠팡'                      → 직매입(제외)
     - 사업자번호가 국내 형식(000-00-00000) 아님 → 해외(제외)  ※ 판매자배송·판매자로켓에 섞인 해외업체 대비

[설치]  pip install playwright pandas openpyxl
[실행]  python coupang_seller_crawler.py --keywords 텀블러 물티슈 --pages 1
        python coupang_seller_crawler.py --keywords --f                      # keywords_coupang.py 전체
        python coupang_seller_crawler.py --keywords 텀블러 --exclude-seller-rocket   # 판매자로켓 제외
        python coupang_seller_crawler.py --keywords 텀블러 --max 5                   # 검색어마다 5건만 시험
        python coupang_seller_crawler.py --keywords --f --delay 3 6            # 대기시간 직접 지정
"""

import argparse
import csv
import os
import random
import re
import time
from pathlib import Path
from urllib.parse import quote

import pandas as pd
from playwright.sync_api import sync_playwright

# ---------------------------------------------------------------- 설정
# 실행 환경 설정 (Docker에서는 환경변수로 바꿔 씀)
BROWSER_CHANNEL = os.getenv("BROWSER_CHANNEL", "chrome")    # chrome | chromium(playwright 내장)
HEADLESS = os.getenv("HEADLESS", "0") == "1"                # 쿠팡은 headless면 대부분 차단
PROFILE_DIR = Path(os.getenv("PROFILE_DIR", "./chrome_profile"))
PROGRESS_CSV = Path("./progress_coupang.csv")   # 중간저장(이어받기용, 실행할 때마다 누적)


def output_path():
    """사용자용 엑셀 파일명: 2026-09-17 082017_coupang.xlsx"""
    return Path(f"./{time.strftime('%Y-%m-%d %H%M%S')}_coupang.xlsx")


def load_keywords(site):
    """keywords_{site}.py 안의 KEYWORDS 리스트를 읽음 (스크립트 폴더 → 현재 폴더 순으로 찾음)"""
    name = f"keywords_{site}.py"
    for p in (Path(__file__).resolve().parent / name, Path.cwd() / name):
        if p.exists():
            ns = {}
            exec(p.read_text(encoding="utf-8"), ns)
            kws = [str(k).strip() for k in ns.get("KEYWORDS", []) if str(k).strip()]
            print(f"[키워드] {p} 에서 {len(kws)}개 읽음")
            return kws
    raise SystemExit(f"{name} 파일을 찾을 수 없습니다. 스크립트와 같은 폴더에 만들어 주세요.")


def resolve_keywords(args, site):
    kws = list(args.keywords or [])
    if args.f:
        kws += load_keywords(site)
    kws = list(dict.fromkeys(kws))                 # 중복 제거(순서 유지)
    if not kws:
        raise SystemExit("검색어가 없습니다. --keywords 텀블러 … 또는 --keywords --f 로 실행하세요.")
    return kws

DELAY = (2, 4)          # 상품 간 기본 대기(초). --delay 로 변경, 차단 신호 시 자동으로 늘어남
SEARCH_URL = "https://www.coupang.com/np/search?component=&q={q}&channel=user&page={p}"

COLUMNS = ["검색어", "검색URL", "배송유형", "상품명", "상품URL", "판매자명", "판매자ID", "판매자상점URL", "상호/대표자",
           "사업장소재지", "이메일", "연락처", "통신판매업신고번호", "사업자번호",
           "국내/해외", "판정", "수집일시", "비고"]

LABELS = {                      # 판매자정보 표 라벨 → 컬럼 (공백 제거 후 비교)
    "상호/대표자": "상호/대표자",
    "사업장소재지": "사업장소재지",
    "e-mail": "이메일",
    "연락처": "연락처",
    "통신판매업신고번호": "통신판매업신고번호",
    "사업자번호": "사업자번호",
}


def nap(a=DELAY[0], b=DELAY[1]):
    time.sleep(random.uniform(a, b))


class Pacer:
    """상품 간 대기 조절: 평소엔 짧게, 차단·오류 신호가 오면 자동으로 늘리고 쉬었다가 다시 줄임"""

    def __init__(self, lo, hi):
        self.lo, self.hi = lo, hi
        self.factor = 1.0          # 대기시간 배수 (1 → 2 → 4 → 8)
        self.ok = 0                # 연속 정상 건수

    def wait(self):
        time.sleep(random.uniform(self.lo, self.hi) * self.factor)
        if random.random() < 0.03:                 # 약 30건에 한 번 15~30초 휴식(일정한 리듬 회피)
            time.sleep(random.uniform(15, 30))

    def success(self):
        self.ok += 1
        if self.factor > 1 and self.ok >= 20:      # 20건 연속 정상이면 대기 배수 절반으로
            self.factor = max(1.0, self.factor / 2)
            self.ok = 0

    def blocked(self):
        self.ok = 0
        self.factor = min(self.factor * 2, 8)
        cool = random.uniform(60, 120)
        print(f"    ⚠ 차단/이상 신호 → {cool:.0f}초 쉬고, 이후 대기시간 ×{self.factor:g}")
        time.sleep(cool)


def wait_products(page, selector, max_scroll=10):
    """상품 링크가 뜰 때까지 기다린 뒤, 개수가 더 안 늘 때까지만 스크롤"""
    try:
        page.wait_for_selector(selector, timeout=15000)
    except Exception:
        return
    last = -1
    for _ in range(max_scroll):
        n = page.locator(selector).count()
        if n == last:
            break
        last = n
        page.mouse.wheel(0, 3000)
        time.sleep(0.4)


# ---------------------------------------------------------------- 1) 검색결과 카드 수집
CARDS_JS = """
() => {
  const out = [];
  document.querySelectorAll('li').forEach(li => {
    const a = li.querySelector('a[href*="/vp/products/"]');
    const img = li.querySelector('img');
    if (!a || !img || li.querySelector('li')) return;
    const srcs = [...li.querySelectorAll('img')].map(i => i.src).join(' ');
    let kind = '판매자배송';
    if (/logo_rocket_filter/.test(srcs)) kind = '로켓배송';
    else if (/logo_jikgu/.test(srcs)) kind = '로켓직구';
    else if (/logo_rocket_merchant/.test(srcs)) kind = '판매자로켓';
    else if (/rocket/.test(srcs)) kind = '기타로켓';
    const u = new URL(a.href);   // 같은 상품도 vendorItemId가 다르면 판매자가 다를 수 있어 유지
    const url = u.origin + u.pathname + '?itemId=' + (u.searchParams.get('itemId') || '')
              + '&vendorItemId=' + (u.searchParams.get('vendorItemId') || '');
    out.push({url, name: img.alt || '', kind,
              ad: /srp_product_ads/.test(a.href)});
  });
  return out;
}
"""


def collect_cards(page, keyword, pages, allowed_kinds):
    cards, seen = [], set()
    for p in range(1, pages + 1):
        page.goto(SEARCH_URL.format(q=quote(keyword), p=p), wait_until="domcontentloaded")
        wait_products(page, 'a[href*="/vp/products/"]')
        found = page.evaluate(CARDS_JS)
        if not found:
            print(f"[검색] '{keyword}' p{p}: 결과 없음(마지막 페이지이거나 차단) → 중단")
            break
        stat = {}
        for c in found:
            stat[c["kind"]] = stat.get(c["kind"], 0) + 1
            if c["kind"] in allowed_kinds and c["url"] not in seen:
                seen.add(c["url"])
                cards.append(c | {"keyword": keyword,
                                  "search_url": SEARCH_URL.format(q=quote(keyword), p=p)})
        print(f"[검색] '{keyword}' p{p}: {stat} → 대상 누적 {len(cards)}")
        if p < pages:
            nap(1.5, 3)
    return cards


# ---------------------------------------------------------------- 2) 상품페이지 판매자정보
SELLER_JS = """
async () => {
  const sleep = ms => new Promise(r => setTimeout(r, ms));
  let tab = null;
  for (let i = 0; i < 20 && !tab; i++) {          // 탭이 그려질 때까지 최대 10초 (고정 대기 대신)
    tab = [...document.querySelectorAll('a,button,li')]
      .find(e => (e.innerText || '').trim() === '배송/교환/반품 안내');
    if (!tab) await sleep(500);
  }
  if (tab) { tab.scrollIntoView(); tab.click(); }
  const loaded = () => /사업자번호|판매자\\s*쿠팡/.test(document.body.innerText);   // 쿠팡 직매입은 사업자번호 표가 없음
  for (let i = 0; i < 20 && !loaded(); i++) await sleep(500);

  const pairs = [];
  document.querySelectorAll('tr').forEach(tr => {
    const c = [...tr.querySelectorAll('th,td')].map(x => x.innerText.trim());
    for (let i = 0; i + 1 < c.length; i += 2) pairs.push([c[i], c[i + 1]]);
  });
  const text = document.body.innerText;
  const shop = document.querySelector('a[href*="shop.coupang.com/vid/"]');
  const m = text.match(/판매자\\s*[:：]\\s*([^\\n]+)/);
  return {
    tabFound: !!tab,
    pairs,
    shopHref: shop ? shop.href : '',
    shopName: shop ? shop.innerText.split('\\n')[0].trim() : '',
    headerSeller: m ? m[1].trim() : '',
    blocked: /Access Denied|비정상적인 접근/.test(text.slice(0, 2000)),
  };
}
"""


def scrape(page, card):
    row = {c: "" for c in COLUMNS}
    row.update({"검색어": card["keyword"], "검색URL": card["search_url"], "배송유형": card["kind"], "상품명": card["name"],
                "상품URL": card["url"], "수집일시": time.strftime("%Y-%m-%d %H:%M:%S")})

    page.goto(card["url"], wait_until="domcontentloaded", timeout=45000)
    r = page.evaluate(SELLER_JS)

    if r["blocked"]:
        row["비고"] = "차단됨"
        return row

    row["판매자명"] = r["shopName"] or r["headerSeller"]
    vid = re.search(r"/vid/([A-Za-z0-9]+)", r["shopHref"])
    row["판매자ID"] = vid.group(1) if vid else ""
    row["판매자상점URL"] = f"https://shop.coupang.com/vid/{vid.group(1)}" if vid else ""

    for label, value in r["pairs"]:
        col = LABELS.get(label.replace(" ", "").replace("\n", ""))
        if col and not row[col]:
            first = value.split("\n")[0] if col == "연락처" else value   # 연락처 뒤 안내문구 제거
            row[col] = re.sub(r"\s+", " ", first).strip()
    if row["사업자번호"]:
        digits = re.sub(r"[\s-]", "", row["사업자번호"])
        if re.fullmatch(r"\d{10}", digits):                         # 국내 사업자등록번호 10자리
            row["사업자번호"] = f"{digits[:3]}-{digits[3:5]}-{digits[5:]}"
            row["국내/해외"] = "국내"
        else:
            row["국내/해외"] = "해외"

    if row["판매자명"].startswith("쿠팡"):
        row["판정"] = "직매입(제외)"
    elif row["국내/해외"] == "해외":
        row["판정"] = "해외(제외)"
    elif row["국내/해외"] == "국내":
        row["판정"] = "중개"
    else:
        row["판정"] = "확인필요"
        row["비고"] = "탭 없음" if not r["tabFound"] else "판매자정보 표 못찾음"
    return row


# ---------------------------------------------------------------- 3) 저장
def done_urls():
    if not PROGRESS_CSV.exists():
        return set()
    with PROGRESS_CSV.open(encoding="utf-8-sig") as f:
        return {r["상품URL"] for r in csv.DictReader(f) if r["비고"] != "차단됨"}


def save_row(row):
    # 예전 버전 progress.csv(컬럼 구성이 다름)가 있으면 백업 후 새로 시작
    if PROGRESS_CSV.exists():
        with PROGRESS_CSV.open(encoding="utf-8-sig") as f:
            if next(csv.reader(f), []) != COLUMNS:
                PROGRESS_CSV.rename(PROGRESS_CSV.with_suffix(f".old_{int(time.time())}.csv"))
    new = not PROGRESS_CSV.exists()
    with PROGRESS_CSV.open("a", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        if new:
            w.writeheader()
        w.writerow(row)


def export_excel(run_map):
    """run_map: {상품URL: {"kw": [검색어…], "search": [검색URL…]}} — 이번 실행에서 검색된 상품"""
    out = output_path()
    df = pd.read_csv(PROGRESS_CSV, encoding="utf-8-sig", dtype=str).fillna("")
    df = df[df["상품URL"].isin(run_map)].copy()   # 이전 실행에서 수집된 상품도 이번 검색에 걸렸으면 포함
    df["검색어"] = df["상품URL"].map(lambda u: "\n".join(run_map[u]["kw"]))
    df["검색URL"] = df["상품URL"].map(lambda u: "\n".join(run_map[u]["search"]))
    df = df.drop_duplicates("상품URL", keep="last")
    med = df[df["판정"] == "중개"].copy()

    # 업체 1곳 = 1행. 그 업체를 수집한 상품URL·검색URL을 모두 붙임
    join = lambda x: "\n".join(dict.fromkeys(p for v in x for p in v.split("\n") if p))
    agg = med.groupby("사업자번호", sort=False).agg(
        노출상품수=("상품URL", "size"),
        검색어=("검색어", join),
        수집상품URL_전체=("상품URL", join),
        검색URL_전체=("검색URL", join),
    ).reset_index()
    first = med.drop_duplicates("사업자번호")[[
        "사업자번호", "판매자명", "판매자ID", "상호/대표자", "사업장소재지", "이메일", "연락처",
        "통신판매업신고번호", "배송유형", "판매자상점URL", "상품명", "상품URL", "수집일시"]]
    first = first.rename(columns={"상품명": "대표상품명", "상품URL": "대표상품URL"})
    uniq = first.merge(agg, on="사업자번호")
    uniq = uniq[["판매자명", "상호/대표자", "사업자번호", "통신판매업신고번호", "사업장소재지",
                 "이메일", "연락처", "판매자ID", "판매자상점URL", "배송유형", "노출상품수",
                 "검색어", "대표상품명", "대표상품URL", "수집상품URL_전체", "검색URL_전체", "수집일시"]]

    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        uniq.to_excel(xw, sheet_name="국내중개업체(중복제거)", index=False)
        df.to_excel(xw, sheet_name="상품별_원자료", index=False)
        df[df["판정"] != "중개"].to_excel(xw, sheet_name="제외·확인필요", index=False)
        style_sheets(xw)
    print(f"[완료] 상품 {len(df)}건 → 국내 중개업체 {len(uniq)}곳 → {out.resolve()}")


def style_sheets(xw):
    """URL 칸은 클릭 가능한 하이퍼링크로, 여러 줄 칸은 줄바꿈 표시"""
    from openpyxl.styles import Alignment, Font
    for ws in xw.book.worksheets:
        headers = [c.value for c in ws[1]]
        ws.freeze_panes = "A2"
        for idx, h in enumerate(headers, 1):
            col = ws.cell(1, idx).column_letter
            ws.column_dimensions[col].width = 45 if "URL" in str(h) else 18
            for cell in ws[col][1:]:
                v = str(cell.value or "")
                if "\n" in v:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                elif v.startswith("http"):
                    cell.hyperlink = v
                    cell.font = Font(color="0563C1", underline="single")


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--keywords", nargs="*", default=[], help="검색어 직접 입력(여러 개 가능)")
    ap.add_argument("--f", action="store_true", help="keywords_coupang.py 의 KEYWORDS 전체 사용")
    ap.add_argument("--pages", type=int, default=1, help="검색어별 페이지 수")
    ap.add_argument("--exclude-seller-rocket", action="store_true",
                    help="판매자로켓(로켓그로스) 상품 제외")
    ap.add_argument("--delay", type=float, nargs=2, default=DELAY, metavar=("최소", "최대"),
                    help="상품 간 대기(초), 기본 2 4")
    ap.add_argument("--max", type=int, default=0, help="검색어별 최대 수집 상품 수(0=제한없음)")
    args = ap.parse_args()
    keywords = resolve_keywords(args, "coupang")

    allowed = {"판매자배송", "판매자로켓", "기타로켓"}     # 로켓배송(직매입)·로켓직구(해외) 제외
    if args.exclude_seller_rocket:
        allowed.discard("판매자로켓")

    pacer = Pacer(*args.delay)
    done = done_urls()
    run_map = {}            # 이번 실행에서 검색된 상품URL → 검색어/검색URL (엑셀용)
    stopped = False

    with sync_playwright() as pw:
        ctx = pw.chromium.launch_persistent_context(
            user_data_dir=str(PROFILE_DIR),
            channel=None if BROWSER_CHANNEL == "chromium" else BROWSER_CHANNEL,
            headless=HEADLESS,
            locale="ko-KR", viewport={"width": 1400, "height": 900},
            args=["--disable-blink-features=AutomationControlled"],
        )
        page = ctx.pages[0] if ctx.pages else ctx.new_page()

        try:
            for k, kw in enumerate(keywords, 1):          # 검색어 하나씩: 검색 → 바로 수집
                print(f"\n===== [{k}/{len(keywords)}] {kw} =====")
                cards = collect_cards(page, kw, args.pages, allowed)
                for c in cards:
                    m = run_map.setdefault(c["url"], {"kw": [], "search": []})
                    m["kw"].append(kw); m["search"].append(c["search_url"])
                todo = [c for c in cards if c["url"] not in done]
                skipped = len(cards) - len(todo)
                if args.max:
                    todo = todo[: args.max]
                print(f"[대상] {len(cards)}건 / 이미 수집 {skipped} / 이번 수집 {len(todo)}")

                streak = 0
                for i, card in enumerate(todo, 1):
                    try:
                        row = scrape(page, card)
                    except Exception as e:
                        row = {c: "" for c in COLUMNS} | {"상품URL": card["url"], "검색어": kw,
                                                          "검색URL": card["search_url"],
                                                          "배송유형": card["kind"], "비고": f"오류: {e!s:.80}"}
                    save_row(row)
                    if row["비고"] != "차단됨":
                        done.add(card["url"])              # 다른 검색어에서 같은 상품 나오면 재수집 안 함
                    print(f"[{i}/{len(todo)}] {row['배송유형']:5} | {row['판정'] or '-':6} | "
                          f"{row['판매자명'] or '-'} | {row['사업자번호'] or '-'} {row['비고']}")
                    if row["비고"] == "차단됨":
                        streak += 1
                        if streak >= 3:
                            print("연속 차단 → 중단. 시간을 두고 다시 실행하면 이어서 진행합니다.")
                            stopped = True
                            break
                        pacer.blocked()
                        continue
                    streak = 0
                    pacer.success()
                    pacer.wait()
                if stopped:
                    break
        except KeyboardInterrupt:
            print("\n사용자 중단(Ctrl+C) → 지금까지 수집한 내용으로 엑셀을 만듭니다.")
        finally:
            ctx.close()

    if PROGRESS_CSV.exists() and run_map:
        export_excel(run_map)


if __name__ == "__main__":
    main()