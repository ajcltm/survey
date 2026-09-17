"""
오픈마켓 체감도조사 - 11번가 입점업체 판매자정보 수집기
------------------------------------------------------------------------
2026-09-17 실제 11번가 페이지 구조를 확인하고 맞춘 버전입니다.
쿠팡 수집기(coupang_seller_crawler.py)와 엑셀 필드·파일 규칙이 같습니다.

[흐름]
  1) 검색결과(search.11st.co.kr)에서 상품 목록 수집
  2) 상품페이지 → '판매자정보 (반품/교환)' 탭 → 판매자명·스토어ID 확인
     ※ 11번가 직매입(슈팅배송, 판매자 '십일번가 주식회사')은 입점업체가 아니므로 제외 (2026-09-17 확인)
  3) '상세정보 확인' 창(보안문자) 열기
       → 사용자가 창에 숫자를 직접 입력하고 [확인]  ※ 보안문자는 자동으로 풀지 않습니다
       → 판매자 상세표(사업자등록번호, 대표번호, e-mail, 영업소재지 …) 추출
  4) 전화번호(고객문의 대표번호)가 없으면 제외 / 사업자번호가 국내 형식이 아니면 제외
  5) 이미 수집한 판매자(같은 스토어ID)의 다른 상품은 보안문자 없이 기존 정보 재사용

[설치]  pip install playwright pandas openpyxl
[실행]  python elevenst_seller_crawler.py --keywords 텀블러 --max 5      # 검색어마다 5건
        python elevenst_seller_crawler.py --keywords --f                  # keywords_11st.py 전체
        python elevenst_seller_crawler.py --keywords --f --delay 2 4       # 대기시간 직접 지정
        python elevenst_seller_crawler.py --keywords 텀블러 물티슈 --pages 2
[결과]  progress_11st.csv                         (중간저장, 누적)
        2026-09-17 082017_11st.xlsx               (사용자용)
"""

import argparse
import csv
import os
import random
import re
import time
from pathlib import Path
from urllib.parse import quote

import pandas as pd
from playwright.sync_api import sync_playwright

# ---------------------------------------------------------------- 설정
BROWSER_CHANNEL = os.getenv("BROWSER_CHANNEL", "chrome")    # chrome | chromium
HEADLESS = os.getenv("HEADLESS", "0") == "1"                # 보안문자 입력이 필요하므로 창 필수
PROFILE_DIR = Path(os.getenv("PROFILE_DIR", "./chrome_profile_11st"))
PROGRESS_CSV = Path("./progress_11st.csv")
DELAY = (1.5, 3)               # 보안문자 없이 넘어간 상품 뒤 대기(초). 보안문자 입력 건은 사람 입력시간이 대기 역할
CAPTCHA_TIMEOUT = 180          # 보안문자 입력 대기(초). 넘기면 해당 상품 건너뜀

SEARCH_URL = "https://search.11st.co.kr/pc/total-search?kwd={q}&tabId=TOTAL_SEARCH&pageNo={p}"
PRODUCT_URL = "https://www.11st.co.kr/products/{no}"
CAPTCHA_URL = "https://m.11st.co.kr/products/mw/pages?area=captcha&prdNo={no}"

# 쿠팡과 동일한 필드
COLUMNS = ["검색어", "검색URL", "배송유형", "상품명", "상품URL", "판매자명", "판매자ID", "판매자상점URL", "상호/대표자",
           "사업장소재지", "이메일", "연락처", "통신판매업신고번호", "사업자번호",
           "국내/해외", "판정", "수집일시", "비고"]

# 보안문자 통과 후 표 라벨(공백 제거) → 컬럼
LABELS = {
    "판매자": "판매자명",
    "상호명/대표자": "상호/대표자",
    "사업자등록번호": "사업자번호",
    "통신판매업신고": "통신판매업신고번호",
    "고객문의대표번호": "연락처",
    "e-mail": "이메일",
    "영업소재지": "사업장소재지",
}
SELLER_FIELDS = ["판매자명", "상호/대표자", "사업장소재지", "이메일", "연락처",
                 "통신판매업신고번호", "사업자번호", "국내/해외", "판정"]
PHONE_RE = re.compile(r"0\d{1,2}-?\d{3,4}-?\d{4}|1\d{3}-?\d{4}")   # 일반/휴대폰/070/1588형


def nap(a=DELAY[0], b=DELAY[1]):
    time.sleep(random.uniform(a, b))


class Pacer:
    """상품 간 대기 조절: 평소엔 짧게, 차단·오류 신호가 오면 자동으로 늘리고 쉬었다가 다시 줄임"""

    def __init__(self, lo, hi):
        self.lo, self.hi = lo, hi
        self.factor = 1.0          # 대기시간 배수 (1 → 2 → 4 → 8)
        self.ok = 0                # 연속 정상 건수

    def wait(self):
        time.sleep(random.uniform(self.lo, self.hi) * self.factor)
        if random.random() < 0.03:                 # 약 30건에 한 번 15~30초 휴식(일정한 리듬 회피)
            time.sleep(random.uniform(15, 30))

    def success(self):
        self.ok += 1
        if self.factor > 1 and self.ok >= 20:      # 20건 연속 정상이면 대기 배수 절반으로
            self.factor = max(1.0, self.factor / 2)
            self.ok = 0

    def blocked(self):
        self.ok = 0
        self.factor = min(self.factor * 2, 8)
        cool = random.uniform(60, 120)
        print(f"    ⚠ 차단/이상 신호 → {cool:.0f}초 쉬고, 이후 대기시간 ×{self.factor:g}")
        time.sleep(cool)


def wait_products(page, selector, max_scroll=10):
    """상품 링크가 뜰 때까지 기다린 뒤, 개수가 더 안 늘 때까지만 스크롤"""
    try:
        page.wait_for_selector(selector, timeout=15000)
    except Exception:
        return
    last = -1
    for _ in range(max_scroll):
        n = page.locator(selector).count()
        if n == last:
            break
        last = n
        page.mouse.wheel(0, 3000)
        time.sleep(0.4)


def output_path():
    """사용자용 엑셀 파일명: 2026-09-17 082017_11st.xlsx"""
    return Path(f"./{time.strftime('%Y-%m-%d %H%M%S')}_11st.xlsx")


def load_keywords(site):
    """keywords_{site}.py 안의 KEYWORDS 리스트를 읽음 (스크립트 폴더 → 현재 폴더 순으로 찾음)"""
    name = f"keywords_{site}.py"
    for p in (Path(__file__).resolve().parent / name, Path.cwd() / name):
        if p.exists():
            ns = {}
            exec(p.read_text(encoding="utf-8"), ns)
            kws = [str(k).strip() for k in ns.get("KEYWORDS", []) if str(k).strip()]
            print(f"[키워드] {p} 에서 {len(kws)}개 읽음")
            return kws
    raise SystemExit(f"{name} 파일을 찾을 수 없습니다. 스크립트와 같은 폴더에 만들어 주세요.")


def resolve_keywords(args, site):
    kws = list(args.keywords or [])
    if args.f:
        kws += load_keywords(site)
    kws = list(dict.fromkeys(kws))                 # 중복 제거(순서 유지)
    if not kws:
        raise SystemExit("검색어가 없습니다. --keywords 텀블러 … 또는 --keywords --f 로 실행하세요.")
    return kws



def beep():
    try:
        import winsound
        winsound.Beep(1000, 300)
    except Exception:
        print("\a", end="")


# ---------------------------------------------------------------- 1) 검색결과
CARDS_JS = """
() => {
  const seen = new Set(), out = [];
  document.querySelectorAll('a[href*="11st.co.kr/products/"]').forEach(a => {
    const m = a.href.match(/\\/products\\/(\\d+)/);
    if (!m || seen.has(m[1])) return;
    const name = (a.innerText || a.title || a.querySelector('img')?.alt || '').replace(/\\s+/g, ' ').trim();
    if (!name) return;                      // 이미지만 있는 중복 링크는 제목 링크에서 잡힘
    seen.add(m[1]);
    out.push({no: m[1], name});
  });
  return out;
}
"""


def collect_cards(page, keyword, pages):
    cards, seen = [], set()
    for p in range(1, pages + 1):
        url = SEARCH_URL.format(q=quote(keyword), p=p)
        page.goto(url, wait_until="domcontentloaded")
        wait_products(page, 'a[href*="11st.co.kr/products/"]')
        found = [c for c in page.evaluate(CARDS_JS) if c["no"] not in seen]
        if not found:
            print(f"[검색] '{keyword}' p{p}: 새 상품 없음 → 중단")
            break
        for c in found:
            seen.add(c["no"])
            cards.append(c | {"keyword": keyword, "search_url": url, "url": PRODUCT_URL.format(no=c["no"])})
        print(f"[검색] '{keyword}' p{p}: {len(found)}개 → 누적 {len(cards)}")
        if p < pages:
            nap(1.5, 3)
    return cards


# ---------------------------------------------------------------- 2) 상품페이지(보안문자 전)
PRODUCT_JS = """
async () => {
  const sleep = ms => new Promise(r => setTimeout(r, ms));
  const findTab = () => document.getElementById('tabMenuDetail4')
    || [...document.querySelectorAll('button')].find(b => /판매자정보/.test(b.innerText));
  const findBtn = () => [...document.querySelectorAll('button')].some(b => b.innerText.trim() === '상세정보 확인');
  let tab = null;
  for (let i = 0; i < 20 && !tab; i++) { tab = findTab(); if (!tab) await sleep(500); }   // 최대 10초
  if (tab) {
    tab.scrollIntoView(); tab.click();
    for (let i = 0; i < 10 && !findBtn(); i++) await sleep(300);                         // 버튼 뜨면 바로 진행
  }
  const store = document.querySelector('a[href*="shop.11st.co.kr/stores/"]');
  const panel = document.getElementById('tabpanelDetail4') || document.body;
  const pairs = {};
  panel.querySelectorAll('tr').forEach(tr => {
    const c = [...tr.querySelectorAll('th,td')].map(x => x.innerText.trim());
    for (let i = 0; i + 1 < c.length; i += 2) pairs[c[i]] = c[i + 1];
  });
  return {
    tabFound: !!tab,
    storeHref: store ? store.href : '',
    seller: pairs['판매자'] || (store ? store.innerText.trim() : ''),
    corp: pairs['상호명'] || pairs['상호명/대표자'] || '',
    hasDetailBtn: findBtn(),
  };
}
"""

# 보안문자 통과 후 표
DETAIL_JS = """
() => {
  const t = document.querySelector('table.c-table');
  if (!t) return null;
  return [...t.querySelectorAll('tr')].map(tr =>
    [tr.querySelector('th')?.innerText.trim() || '', tr.querySelector('td')?.innerText.trim() || '']);
}
"""


def wait_captcha(cap, no):
    """보안문자 창을 열고 사용자가 입력할 때까지 대기 → 상세표 반환(시간초과 시 None)"""
    cap.goto(CAPTCHA_URL.format(no=no), wait_until="domcontentloaded")
    cap.bring_to_front()
    beep()
    print(f"    ▶ 브라우저 창에 보안문자 숫자를 입력하고 [확인]을 눌러주세요 (최대 {CAPTCHA_TIMEOUT}초)")
    end = time.time() + CAPTCHA_TIMEOUT
    while time.time() < end:
        try:
            rows = cap.evaluate(DETAIL_JS)
            if rows:
                return rows
        except Exception:
            pass                        # 입력 후 페이지 갱신 중
        time.sleep(0.5)
    return None


def apply_detail(row, rows):
    for label, value in rows:
        col = LABELS.get(label.replace(" ", ""))
        if col:
            row[col] = re.sub(r"\s+", " ", value).strip()
    row["상호/대표자"] = row["상호/대표자"].replace(" | ", " / ")

    m = PHONE_RE.search(row["연락처"])           # '070-4140-1248(유료)' → '070-4140-1248'
    row["연락처"] = m.group(0) if m else ""

    if row["사업자번호"]:
        digits = re.sub(r"[\s-]", "", row["사업자번호"])
        if re.fullmatch(r"\d{10}", digits):
            row["사업자번호"] = f"{digits[:3]}-{digits[3:5]}-{digits[5:]}"
            row["국내/해외"] = "국내"
        else:
            row["국내/해외"] = "해외"

    if not row["연락처"]:
        row["판정"] = "전화번호 없음(제외)"
    elif row["국내/해외"] == "해외":
        row["판정"] = "해외(제외)"
    else:
        row["판정"] = "중개"


def scrape(page, cap, card, seller_cache):
    row = {c: "" for c in COLUMNS}
    row.update({"검색어": card["keyword"], "검색URL": card["search_url"], "상품명": card["name"],
                "상품URL": card["url"], "수집일시": time.strftime("%Y-%m-%d %H:%M:%S")})

    page.goto(card["url"], wait_until="domcontentloaded", timeout=45000)
    info = page.evaluate(PRODUCT_JS)

    sid = re.search(r"/stores/(\d+)", info["storeHref"])
    row["판매자ID"] = sid.group(1) if sid else ""
    row["판매자상점URL"] = f"https://shop.11st.co.kr/stores/{row['판매자ID']}" if sid else ""
    row["판매자명"] = info["seller"]

    # 같은 판매자를 이미 수집했으면 보안문자 생략
    if row["판매자ID"] and row["판매자ID"] in seller_cache:
        row.update(seller_cache[row["판매자ID"]])
        row["비고"] = "기존 수집 판매자 정보 재사용"
        return row

    # 11번가 직매입(슈팅배송 등): 판매자가 십일번가(주), 보안문자 버튼 없이 정보가 바로 보임 → 제외
    if "십일번가" in info["corp"] or info["seller"] == "슈팅배송":
        row["판매자명"] = row["판매자명"] or info["seller"]
        row["상호/대표자"] = info["corp"].replace(" | ", " / ")
        row["판정"] = "직매입(제외)"
        return row

    if not info["hasDetailBtn"]:
        row["판정"] = "확인필요"
        row["비고"] = "상세정보 확인 버튼 없음"
        return row

    rows = wait_captcha(cap, card["no"])
    page.bring_to_front()
    if rows is None:
        row["판정"] = "보안문자 미입력"
        row["비고"] = "시간초과로 건너뜀(다음 실행 때 재시도)"
        return row

    apply_detail(row, rows)
    row["_captcha"] = True
    if row["판매자ID"]:
        seller_cache[row["판매자ID"]] = {k: row[k] for k in SELLER_FIELDS}
    return row


# ---------------------------------------------------------------- 3) 저장 (쿠팡과 동일)
def read_progress():
    if not PROGRESS_CSV.exists():
        return []
    with PROGRESS_CSV.open(encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def save_row(row):
    # 예전 버전 progress(컬럼 구성이 다름)가 있으면 백업 후 새로 시작
    if PROGRESS_CSV.exists():
        with PROGRESS_CSV.open(encoding="utf-8-sig") as f:
            if next(csv.reader(f), []) != COLUMNS:
                PROGRESS_CSV.rename(PROGRESS_CSV.with_suffix(f".old_{int(time.time())}.csv"))
    new = not PROGRESS_CSV.exists()
    with PROGRESS_CSV.open("a", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        if new:
            w.writeheader()
        w.writerow({k: row.get(k, "") for k in COLUMNS})


def export_excel(run_map):
    """run_map: {상품URL: {"kw": [검색어…], "search": [검색URL…]}} — 이번 실행에서 검색된 상품"""
    out = output_path()
    df = pd.read_csv(PROGRESS_CSV, encoding="utf-8-sig", dtype=str).fillna("")
    df = df[df["상품URL"].isin(run_map)].copy()   # 이전 실행에서 수집된 상품도 이번 검색에 걸렸으면 포함
    df["검색어"] = df["상품URL"].map(lambda u: "\n".join(run_map[u]["kw"]))
    df["검색URL"] = df["상품URL"].map(lambda u: "\n".join(run_map[u]["search"]))
    df = df.drop_duplicates("상품URL", keep="last")
    med = df[df["판정"] == "중개"].copy()

    join = lambda x: "\n".join(dict.fromkeys(p for v in x for p in v.split("\n") if p))
    agg = med.groupby("사업자번호", sort=False).agg(
        노출상품수=("상품URL", "size"),
        검색어=("검색어", join),
        수집상품URL_전체=("상품URL", join),
        검색URL_전체=("검색URL", join),
    ).reset_index()
    first = med.drop_duplicates("사업자번호")[[
        "사업자번호", "판매자명", "판매자ID", "상호/대표자", "사업장소재지", "이메일", "연락처",
        "통신판매업신고번호", "배송유형", "판매자상점URL", "상품명", "상품URL", "수집일시"]]
    first = first.rename(columns={"상품명": "대표상품명", "상품URL": "대표상품URL"})
    uniq = first.merge(agg, on="사업자번호")
    uniq = uniq[["판매자명", "상호/대표자", "사업자번호", "통신판매업신고번호", "사업장소재지",
                 "이메일", "연락처", "판매자ID", "판매자상점URL", "배송유형", "노출상품수",
                 "검색어", "대표상품명", "대표상품URL", "수집상품URL_전체", "검색URL_전체", "수집일시"]]

    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        uniq.to_excel(xw, sheet_name="국내중개업체(중복제거)", index=False)
        df.to_excel(xw, sheet_name="상품별_원자료", index=False)
        df[df["판정"] != "중개"].to_excel(xw, sheet_name="제외·확인필요", index=False)
        style_sheets(xw)
    print(f"[완료] 상품 {len(df)}건 → 국내 중개업체 {len(uniq)}곳 → {out.resolve()}")


def style_sheets(xw):
    from openpyxl.styles import Alignment, Font
    for ws in xw.book.worksheets:
        headers = [c.value for c in ws[1]]
        ws.freeze_panes = "A2"
        for idx, h in enumerate(headers, 1):
            col = ws.cell(1, idx).column_letter
            ws.column_dimensions[col].width = 45 if "URL" in str(h) else 18
            for cell in ws[col][1:]:
                v = str(cell.value or "")
                if "\n" in v:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                elif v.startswith("http"):
                    cell.hyperlink = v
                    cell.font = Font(color="0563C1", underline="single")


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--keywords", nargs="*", default=[], help="검색어 직접 입력(여러 개 가능)")
    ap.add_argument("--f", action="store_true", help="keywords_11st.py 의 KEYWORDS 전체 사용")
    ap.add_argument("--pages", type=int, default=1, help="검색어별 페이지 수")
    ap.add_argument("--delay", type=float, nargs=2, default=DELAY, metavar=("최소", "최대"),
                    help="보안문자 없이 넘어간 상품 뒤 대기(초), 기본 1.5 3")
    ap.add_argument("--max", type=int, default=0, help="검색어별 최대 수집 상품 수(0=제한없음)")
    args = ap.parse_args()
    keywords = resolve_keywords(args, "11st")

    prev = read_progress()
    # 보안문자 미입력·확인필요 건은 다음 실행 때 다시 시도
    retry = ("보안문자 미입력", "확인필요", "")
    done = {r["상품URL"] for r in prev if r["판정"] not in retry}
    seller_cache = {r["판매자ID"]: {k: r[k] for k in SELLER_FIELDS}
                    for r in prev if r["판매자ID"] and r["판정"] in ("중개", "전화번호 없음(제외)", "해외(제외)")}
    run_map = {}
    pacer = Pacer(*args.delay)

    with sync_playwright() as pw:
        ctx = pw.chromium.launch_persistent_context(
            user_data_dir=str(PROFILE_DIR),
            channel=None if BROWSER_CHANNEL == "chromium" else BROWSER_CHANNEL,
            headless=HEADLESS, locale="ko-KR", viewport={"width": 1300, "height": 900},
            args=["--disable-blink-features=AutomationControlled"],
        )
        page = ctx.pages[0] if ctx.pages else ctx.new_page()
        cap = ctx.new_page()                    # 보안문자 전용 탭
        page.bring_to_front()

        try:
            for k, kw in enumerate(keywords, 1):          # 검색어 하나씩: 검색 → 바로 수집
                print(f"\n===== [{k}/{len(keywords)}] {kw} =====")
                cards = collect_cards(page, kw, args.pages)
                for c in cards:
                    m = run_map.setdefault(c["url"], {"kw": [], "search": []})
                    m["kw"].append(kw); m["search"].append(c["search_url"])
                todo = [c for c in cards if c["url"] not in done]
                skipped = len(cards) - len(todo)
                if args.max:
                    todo = todo[: args.max]
                print(f"[대상] {len(cards)}건 / 이미 수집 {skipped} / 이번 수집 {len(todo)}"
                      f" / 확보 판매자 {len(seller_cache)}곳")

                bad = 0
                for i, card in enumerate(todo, 1):
                    print(f"[{i}/{len(todo)}] {card['name'][:40]}")
                    try:
                        row = scrape(page, cap, card, seller_cache)
                    except Exception as e:
                        row = {c: "" for c in COLUMNS} | {"상품URL": card["url"], "검색어": kw,
                                                          "검색URL": card["search_url"], "판정": "확인필요",
                                                          "비고": f"오류: {e!s:.80}"}
                    save_row(row)
                    if row["판정"] not in retry:
                        done.add(card["url"])
                    print(f"    → {row['판정']:12} | {row['판매자명'] or '-'} | {row['연락처'] or '-'} | {row['비고']}")
                    if row["판정"] == "확인필요":             # 탭·버튼을 못 찾음 = 차단/오류 가능성
                        bad += 1
                        if bad >= 3:
                            pacer.blocked()
                            bad = 0
                        continue
                    bad = 0
                    pacer.success()
                    if row.get("_captcha"):                  # 사람이 방금 입력함 → 추가 대기 불필요
                        continue
                    pacer.wait()
        except KeyboardInterrupt:
            print("\n사용자 중단(Ctrl+C) → 지금까지 수집한 내용으로 엑셀을 만듭니다.")
        finally:
            ctx.close()

    if PROGRESS_CSV.exists() and run_map:
        export_excel(run_map)


if __name__ == "__main__":
    main()